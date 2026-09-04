import io
import os
import re
import tempfile
import logging
from typing import List, Tuple, Optional
from PIL import Image
import fitz  # PyMuPDF
from pypdf import PdfReader, PdfWriter
from docx import Document
from pdf2docx import Converter

from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors

logger = logging.getLogger(__name__)

MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024  # 25 MB


class FileConverterService:
    """
    Motor determinístico de alta fidelidade para conversão e manipulação de arquivos.
    Cumpre rigorosamente as regras P-0901 a P-0906 de segurança, isolamento e performance.
    """

    @staticmethod
    def validate_file_size(file_bytes: bytes, max_mb: int = 25) -> None:
        """Valida se o arquivo está dentro do limite máximo em megabytes (P-0901)."""
        limit = max_mb * 1024 * 1024
        if len(file_bytes) > limit:
            size_mb = len(file_bytes) / (1024 * 1024)
            raise ValueError(
                f"Arquivo de {size_mb:.2f} MB excede o limite maximo permitido de {max_mb} MB."
            )

    @staticmethod
    def images_to_pdf(images_bytes_list: List[bytes]) -> bytes:
        """
        Combina uma ou mais imagens (PNG, JPG, JPEG, WEBP, BMP) em um único arquivo PDF.
        """
        if not images_bytes_list:
            raise ValueError("Pelo menos uma imagem deve ser fornecida.")

        pil_images = []
        for img_bytes in images_bytes_list:
            FileConverterService.validate_file_size(img_bytes)
            try:
                img = Image.open(io.BytesIO(img_bytes))
                # Converte para RGB caso seja RGBA ou Paleta, preservando fundo branco
                if img.mode in ("RGBA", "LA", "P"):
                    background = Image.new("RGB", img.size, (255, 255, 255))
                    if img.mode == "P":
                        img = img.convert("RGBA")
                    background.paste(img, mask=img.split()[-1] if img.mode == "RGBA" else None)
                    pil_images.append(background)
                elif img.mode != "RGB":
                    pil_images.append(img.convert("RGB"))
                else:
                    pil_images.append(img)
            except Exception as e:
                raise ValueError(f"Falha ao carregar imagem: {e}")

        output_buf = io.BytesIO()
        first_img = pil_images[0]
        remaining = pil_images[1:] if len(pil_images) > 1 else []

        first_img.save(
            output_buf,
            format="PDF",
            save_all=True,
            append_images=remaining
        )
        return output_buf.getvalue()

    @staticmethod
    def convert_image(image_bytes: bytes, target_format: str = "JPEG", quality: int = 85) -> bytes:
        """
        Transcodifica imagem entre formatos (PNG, JPEG, WEBP, BMP).
        """
        FileConverterService.validate_file_size(image_bytes)
        norm_format = target_format.upper()
        if norm_format in ("JPG", "JPE"):
            norm_format = "JPEG"

        try:
            img = Image.open(io.BytesIO(image_bytes))
        except Exception as e:
            raise ValueError(f"Imagem invalida ou corrompida: {e}")

        if norm_format == "JPEG" and img.mode in ("RGBA", "LA", "P"):
            background = Image.new("RGB", img.size, (255, 255, 255))
            if img.mode == "P":
                img = img.convert("RGBA")
            background.paste(img, mask=img.split()[-1] if img.mode == "RGBA" else None)
            img = background
        elif norm_format != "PNG" and img.mode not in ("RGB", "L"):
            img = img.convert("RGB")

        output_buf = io.BytesIO()
        save_kwargs = {}
        if norm_format in ("JPEG", "WEBP"):
            save_kwargs["quality"] = quality

        img.save(output_buf, format=norm_format, **save_kwargs)
        return output_buf.getvalue()

    @staticmethod
    def merge_pdfs(pdf_bytes_list: List[bytes]) -> bytes:
        """
        Funde dois ou mais arquivos PDF em um único documento mantendo a ordem sequencial.
        """
        if len(pdf_bytes_list) < 2:
            raise ValueError("Necessario pelo menos 2 arquivos PDF para a fusao.")

        writer = PdfWriter()
        total_size = sum(len(b) for b in pdf_bytes_list)
        if total_size > 50 * 1024 * 1024:
            raise ValueError("Tamanho total dos PDFs excede o limite de 50 MB para juncao.")

        for idx, pdf_bytes in enumerate(pdf_bytes_list):
            FileConverterService.validate_file_size(pdf_bytes)
            try:
                reader = PdfReader(io.BytesIO(pdf_bytes))
                if reader.is_encrypted:
                    raise ValueError(f"O PDF {idx + 1} esta protegido por senha. Remova a senha antes.")
                for page in reader.pages:
                    writer.add_page(page)
            except Exception as e:
                if "protegido por senha" in str(e):
                    raise
                raise ValueError(f"PDF {idx + 1} invalido ou corrompido: {e}")

        output_buf = io.BytesIO()
        writer.write(output_buf)
        return output_buf.getvalue()

    @staticmethod
    def split_pdf(pdf_bytes: bytes, pages_range: str) -> bytes:
        """
        Extrai páginas selecionadas de um PDF com base no intervalo (ex: '1-3, 5, 7-9').
        """
        FileConverterService.validate_file_size(pdf_bytes)
        try:
            reader = PdfReader(io.BytesIO(pdf_bytes))
            if reader.is_encrypted:
                raise ValueError("O documento PDF esta protegido por senha.")
        except Exception as e:
            raise ValueError(f"PDF invalido ou corrompido: {e}")

        total_pages = len(reader.pages)
        if total_pages == 0:
            raise ValueError("O documento PDF nao contem paginas.")

        # Parseia o intervalo de páginas (1-indexado)
        selected_pages = set()
        parts = [p.strip() for p in pages_range.split(",") if p.strip()]
        if not parts:
            raise ValueError("Nenhum intervalo de paginas valido foi informado.")

        for part in parts:
            if "-" in part:
                tokens = part.split("-")
                if len(tokens) != 2 or not tokens[0].isdigit() or not tokens[1].isdigit():
                    raise ValueError(f"Intervalo invalido: '{part}'")
                start, end = int(tokens[0]), int(tokens[1])
                if start > end:
                    start, end = end, start
                for page_num in range(start, end + 1):
                    selected_pages.add(page_num)
            elif part.isdigit():
                selected_pages.add(int(part))
            else:
                raise ValueError(f"Numero de pagina invalido: '{part}'")

        # Valida limites
        sorted_pages = sorted(list(selected_pages))
        for page_num in sorted_pages:
            if page_num < 1 or page_num > total_pages:
                raise ValueError(
                    f"Pagina {page_num} fora dos limites do documento (o PDF contem {total_pages} paginas)."
                )

        writer = PdfWriter()
        for page_num in sorted_pages:
            writer.add_page(reader.pages[page_num - 1])

        output_buf = io.BytesIO()
        writer.write(output_buf)
        return output_buf.getvalue()

    @staticmethod
    def pdf_to_text(pdf_bytes: bytes) -> str:
        """
        Extrai todo o texto legível de um arquivo PDF.
        """
        FileConverterService.validate_file_size(pdf_bytes)
        try:
            reader = PdfReader(io.BytesIO(pdf_bytes))
            if reader.is_encrypted:
                raise ValueError("O documento PDF esta protegido por senha.")
        except Exception as e:
            raise ValueError(f"PDF invalido ou corrompido: {e}")

        text_pages = []
        for i, page in enumerate(reader.pages):
            txt = page.extract_text() or ""
            if txt.strip():
                text_pages.append(f"--- Pagina {i + 1} ---\n{txt.strip()}")

        if not text_pages:
            return "Nenhum texto legivel encontrado no documento PDF (possivelmente e um documento digitalizado/imagem)."
        return "\n\n".join(text_pages)

    @staticmethod
    def docx_to_pdf(docx_bytes: bytes) -> bytes:
        """
        Converte documento Word (.docx) em PDF limpo e formatado via ReportLab.
        """
        FileConverterService.validate_file_size(docx_bytes)
        try:
            doc = Document(io.BytesIO(docx_bytes))
        except Exception as e:
            raise ValueError(f"Arquivo DOCX invalido ou corrompido: {e}")

        pdf_buf = io.BytesIO()
        pdf_doc = SimpleDocTemplate(
            pdf_buf,
            pagesize=A4,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()
        normal_style = styles["Normal"]
        h1_style = styles["Heading1"]
        h2_style = styles["Heading2"]
        h3_style = styles["Heading3"]

        story = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                story.append(Spacer(1, 8))
                continue

            # Sanitiza caracteres HTML que o ReportLab interpreta como tags
            safe_text = (
                text.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )

            p_style = normal_style
            if p.style.name.startswith("Heading 1"):
                p_style = h1_style
            elif p.style.name.startswith("Heading 2"):
                p_style = h2_style
            elif p.style.name.startswith("Heading 3"):
                p_style = h3_style

            story.append(Paragraph(safe_text, p_style))
            story.append(Spacer(1, 6))

        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    cell_text = cell.text.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    row_data.append(Paragraph(cell_text, normal_style))
                table_data.append(row_data)

            if table_data:
                t = Table(table_data)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.whitesmoke),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ]))
                story.append(t)
                story.append(Spacer(1, 10))

        if not story:
            story.append(Paragraph("(Documento vazio)", normal_style))

        pdf_doc.build(story)
        return pdf_buf.getvalue()

    @staticmethod
    def pdf_to_docx(pdf_bytes: bytes) -> bytes:
        """
        Converte documento PDF em Word (.docx) preservando estrutura e tabelas (pdf2docx).
        Garante isolamento temporário total (P-0903).
        """
        FileConverterService.validate_file_size(pdf_bytes)

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = os.path.join(temp_dir, "input.pdf")
            docx_path = os.path.join(temp_dir, "output.docx")

            with open(pdf_path, "wb") as f:
                f.write(pdf_bytes)

            try:
                cv = Converter(pdf_path)
                cv.convert(docx_path, start=0, end=None)
                cv.close()
            except Exception as e:
                raise ValueError(f"Falha ao converter PDF para DOCX: {e}")

            if not os.path.exists(docx_path):
                raise ValueError("Nao foi possivel gerar o arquivo DOCX de saida.")

            with open(docx_path, "rb") as f:
                return f.read()

    @staticmethod
    def pdf_to_images(pdf_bytes: bytes, dpi: int = 150) -> List[Tuple[str, bytes]]:
        """
        Renderiza páginas de um PDF como imagens PNG em alta resolução (PyMuPDF).
        Retorna lista de tuplas: (nome_arquivo, bytes_imagem).
        """
        FileConverterService.validate_file_size(pdf_bytes)
        try:
            doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            if doc.is_encrypted:
                raise ValueError("O documento PDF esta protegido por senha.")
        except Exception as e:
            raise ValueError(f"PDF invalido ou corrompido: {e}")

        result_images = []
        try:
            for i, page in enumerate(doc):
                pix = page.get_pixmap(dpi=dpi)
                img_data = pix.tobytes("png")
                filename = f"pagina_{i + 1}.png"
                result_images.append((filename, img_data))
        finally:
            doc.close()

        return result_images
