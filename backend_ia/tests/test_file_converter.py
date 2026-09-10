import io
import pytest
from PIL import Image
from pypdf import PdfReader, PdfWriter
from docx import Document
from reportlab.pdfgen import canvas

from services.file_converter_service import FileConverterService
from repositories.file_conversion_repository import FileConversionRepository
from services.tools.file_converter_tool import FileConverterTool


def _create_sample_image(format="PNG", color=(255, 0, 0), size=(100, 100)) -> bytes:
    buf = io.BytesIO()
    img = Image.new("RGB", size, color=color)
    img.save(buf, format=format)
    return buf.getvalue()


def _create_sample_pdf(pages_count: int = 1, text: str = "Documento Teste") -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i in range(pages_count):
        c.drawString(100, 700, f"{text} - Pagina {i + 1}")
        c.showPage()
    c.save()
    return buf.getvalue()


def _create_sample_docx(text: str = "Conteudo do documento Word de teste") -> bytes:
    doc = Document()
    doc.add_heading("Titulo do Documento", level=1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestFileConverterService:
    def test_images_to_pdf_single_image(self):
        img_bytes = _create_sample_image("PNG")
        pdf_bytes = FileConverterService.images_to_pdf([img_bytes])
        assert pdf_bytes.startswith(b"%PDF-")
        reader = PdfReader(io.BytesIO(pdf_bytes))
        assert len(reader.pages) == 1

    def test_images_to_pdf_multiple_images(self):
        img1 = _create_sample_image("PNG", color=(255, 0, 0))
        img2 = _create_sample_image("JPEG", color=(0, 255, 0))
        pdf_bytes = FileConverterService.images_to_pdf([img1, img2])
        assert pdf_bytes.startswith(b"%PDF-")
        reader = PdfReader(io.BytesIO(pdf_bytes))
        assert len(reader.pages) == 2

    def test_images_to_pdf_empty_list_raises(self):
        with pytest.raises(ValueError, match="Pelo menos uma imagem"):
            FileConverterService.images_to_pdf([])

    def test_convert_image_png_to_jpeg(self):
        png_bytes = _create_sample_image("PNG")
        jpeg_bytes = FileConverterService.convert_image(png_bytes, target_format="JPEG", quality=90)
        img = Image.open(io.BytesIO(jpeg_bytes))
        assert img.format == "JPEG"

    def test_convert_image_jpeg_to_webp(self):
        jpeg_bytes = _create_sample_image("JPEG")
        webp_bytes = FileConverterService.convert_image(jpeg_bytes, target_format="WEBP")
        img = Image.open(io.BytesIO(webp_bytes))
        assert img.format == "WEBP"

    def test_merge_pdfs_success(self):
        pdf1 = _create_sample_pdf(pages_count=2, text="Doc 1")
        pdf2 = _create_sample_pdf(pages_count=3, text="Doc 2")
        merged = FileConverterService.merge_pdfs([pdf1, pdf2])
        assert merged.startswith(b"%PDF-")
        reader = PdfReader(io.BytesIO(merged))
        assert len(reader.pages) == 5

    def test_merge_pdfs_less_than_two_raises(self):
        pdf1 = _create_sample_pdf(pages_count=1)
        with pytest.raises(ValueError, match="Necessario pelo menos 2 arquivos"):
            FileConverterService.merge_pdfs([pdf1])

    def test_split_pdf_page_range(self):
        pdf = _create_sample_pdf(pages_count=5)
        split = FileConverterService.split_pdf(pdf, pages_range="1-2, 4")
        assert split.startswith(b"%PDF-")
        reader = PdfReader(io.BytesIO(split))
        assert len(reader.pages) == 3

    def test_split_pdf_out_of_bounds_raises(self):
        pdf = _create_sample_pdf(pages_count=3)
        with pytest.raises(ValueError, match="fora dos limites"):
            FileConverterService.split_pdf(pdf, pages_range="10")

    def test_pdf_to_text_extraction(self):
        pdf = _create_sample_pdf(pages_count=1, text="PalavraChaveSecreta")
        extracted_text = FileConverterService.pdf_to_text(pdf)
        assert "PalavraChaveSecreta" in extracted_text

    def test_docx_to_pdf_conversion(self):
        docx_bytes = _create_sample_docx(text="Texto de teste para conversao DOCX para PDF")
        pdf_bytes = FileConverterService.docx_to_pdf(docx_bytes)
        assert pdf_bytes.startswith(b"%PDF-")
        reader = PdfReader(io.BytesIO(pdf_bytes))
        assert len(reader.pages) >= 1

    def test_pdf_to_docx_conversion(self):
        pdf = _create_sample_pdf(pages_count=1, text="Texto para ser convertido em docx")
        docx_bytes = FileConverterService.pdf_to_docx(pdf)
        assert docx_bytes.startswith(b"PK")  # ZIP header de DOCX
        doc = Document(io.BytesIO(docx_bytes))
        full_text = "\n".join([p.text for p in doc.paragraphs])
        assert len(full_text.strip()) > 0

    def test_pdf_to_images_conversion(self):
        pdf = _create_sample_pdf(pages_count=2, text="Pagina para renderizar")
        images = FileConverterService.pdf_to_images(pdf, dpi=150)
        assert len(images) == 2
        for filename, img_data in images:
            assert filename.endswith(".png")
            img = Image.open(io.BytesIO(img_data))
            assert img.format == "PNG"

    def test_file_size_exceeded_raises(self):
        # Cria payload fictício maior que 25MB (apenas simulando o tamanho)
        oversized = b"0" * (26 * 1024 * 1024)
        with pytest.raises(ValueError, match="excede o limite maximo"):
            FileConverterService.validate_file_size(oversized, max_mb=25)


class TestFileConversionRepository:
    def test_save_and_retrieve_conversions(self):
        FileConversionRepository.clear_in_memory()
        rec_id = FileConversionRepository.save_conversion(
            user_id="daniel",
            conversion_type="pdf_to_docx",
            source_format="pdf",
            target_format="docx",
            file_size_bytes=1024,
            output_size_bytes=2048,
            status="success",
            execution_time_ms=120
        )
        assert rec_id is not None

        # Isolamento de usuário
        daniel_history = FileConversionRepository.get_user_conversions("daniel")
        assert len(daniel_history) >= 1
        assert daniel_history[0]["conversion_type"] == "pdf_to_docx"

        lari_history = FileConversionRepository.get_user_conversions("lari")
        assert len(lari_history) == 0


class TestFileConverterTool:
    def test_tool_listar_formatos(self):
        res = FileConverterTool.execute(acao="listar_formatos")
        assert "PDF para Word" in res
        assert "Juntar PDFs" in res
        assert "25 MB" in res

    def test_tool_instrucoes(self):
        res = FileConverterTool.execute(acao="instrucoes", tipo_conversao="pdf_para_word")
        assert "PDF" in res
        assert "Word" in res


class TestFilesRouter:
    def test_get_supported_formats(self):
        from fastapi.testclient import TestClient
        from main import app
        client = TestClient(app)
        resp = client.get("/api/files/supported-formats")
        assert resp.status_code == 200
        data = resp.json()
        assert "supported_conversions" in data
        assert data["max_file_size_mb"] == 25

    def test_convert_images_to_pdf_endpoint(self):
        from fastapi.testclient import TestClient
        from main import app
        client = TestClient(app)
        img_bytes = _create_sample_image("PNG")

        files = {
            "file": ("foto.png", img_bytes, "image/png")
        }
        data = {
            "conversion_type": "img_to_pdf"
        }
        resp = client.post("/api/files/convert", files=files, data=data)
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/pdf"
        assert resp.content.startswith(b"%PDF-")

    def test_convert_invalid_type_endpoint(self):
        from fastapi.testclient import TestClient
        from main import app
        client = TestClient(app)
        files = {
            "file": ("teste.txt", b"conteudo", "text/plain")
        }
        data = {
            "conversion_type": "tipo_inexistente"
        }
        resp = client.post("/api/files/convert", files=files, data=data)
        assert resp.status_code == 400


class TestWhatsAppServiceDocument:
    def test_send_document_success(self, monkeypatch):
        import requests
        from services.whatsapp_service import WhatsAppService

        class MockResponse:
            def raise_for_status(self):
                pass
            def json(self):
                return {"status": "SUCCESS", "message": "Document sent"}

        monkeypatch.setattr(requests, "post", lambda *args, **kwargs: MockResponse())
        resp = WhatsAppService.send_document(
            remote_jid="5511999999999@s.whatsapp.net",
            file_bytes=b"%PDF-sample",
            filename="teste.pdf",
            mime_type="application/pdf",
            caption="Segue seu arquivo"
        )
        assert resp is not None
        assert resp.get("status") == "SUCCESS"

